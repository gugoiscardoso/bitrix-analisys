#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Relatório executivo do suporte NÃO fiscal — o espelho do relatorio_executivo.

    python pipeline/gerar_executivo_naofiscal.py --de 2026-05-01 --ate 2026-08-31

Duas abas, não uma. As taxonomias não-fiscais de chamado (11 temas) e de conversa
(12 temas) foram derivadas separadas e não compartilham um único nome de tema, então
somar as duas contagens numa linha só produziria um número sem referente. Cada aba
usa também as métricas que existem no seu canal: chamado tem ciclo em dias e tempo de
desenvolvimento, conversa tem duração e volume de mensagens.

Régua de recência idêntica à do executivo fiscal (importada de gerar_relatorio, não
recopiada) para que "Ativa/Atenção/Antiga" queira dizer a mesma coisa nos dois.

RESSALVA que a Legenda repete: estas duas taxonomias nunca passaram por auditoria
cega. Os 6,7% de erro medidos em 31/08/2026 valem só para a fiscal.
"""
from __future__ import annotations
import argparse
import collections
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gerar_relatorio import (BORDA, CALOR, HDR, HF, TAG_RE, calor_de, dias,  # noqa: E402
                             ler, media_fmt)

RAIZ = Path(__file__).resolve().parent.parent
STORE = RAIZ / "data" / "store"
REPORT = RAIZ / "report"

SEM_CLS = "(sem classificação)"


def carregar(de: str, ate: str):
    """Junta base histórica e os dois caches não-fiscais, separados por tipo."""
    cache = {"chamado": ler(STORE / "classificacao_nao_fiscal.json"),
             "conversa": ler(STORE / "classificacao_nao_fiscal_conv.json")}
    out = {"chamado": [], "conversa": []}
    with (STORE / "base_historica.jsonl").open(encoding="utf-8") as fh:
        for linha in fh:
            r = json.loads(linha)
            if r["fiscal"] or not (de <= r["data"] <= ate):
                continue
            c = cache[r["tipo"]].get(str(r["id"]))
            r["tema"] = c["tema"] if c else SEM_CLS
            out[r["tipo"]].append(r)
    return out


def cabecalho(ws, cols: list[str], larguras: list[int]):
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = HDR, HF
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def formatar(ws, i: int, ncols: int, heat: str, pct: int, centro, calor):
    for c in centro:
        ws.cell(row=i, column=c).alignment = Alignment(horizontal="center", vertical="top")
    for c in calor:
        ws.cell(row=i, column=c).fill, ws.cell(row=i, column=c).font = CALOR[heat]
        ws.cell(row=i, column=c).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=3).font = Font(
        name="Consolas", size=10, bold=pct >= 20,
        color="1F4E79" if pct >= 20 else ("3D83B8" if pct >= 10 else "8FB4D0"))
    for c in range(1, ncols + 1):
        ws.cell(row=i, column=c).border = BORDA
    ws.row_dimensions[i].height = 30


def barra(pct: int) -> str:
    return "█" * min(max(1, round(pct / 4)), 12) + f"  {pct}%"


def aba_chamados(wb: Workbook, regs: list[dict], ate: str) -> int:
    ws = wb.create_sheet("Chamados")
    COLS = ["Tema", "Chamados", "Peso", "Categoria predominante",
            "Tempo total do chamado", "Tempo tech (desenvolvimento)",
            "Última incidência", "Recência"]
    cabecalho(ws, COLS, [64, 11, 15, 24, 19, 19, 14, 11])

    ag = collections.defaultdict(lambda: {"n": 0, "datas": [], "total": [], "tech": [],
                                          "tags": collections.Counter()})
    for r in regs:
        g = ag[r["tema"]]
        g["n"] += 1
        g["datas"].append(r["data"])
        m = TAG_RE.match(r.get("titulo") or "")
        g["tags"][m.group(1).strip() if m else "(sem tag)"] += 1
        if r["status"] == "concluída":
            if (d := dias(r.get("criado_em"), r.get("fechado_em"))) is not None:
                g["total"].append(d)
            if (d := dias(r.get("inicio_dev"), r.get("alterado_em"))) is not None:
                g["tech"].append(d)

    for tema, g in sorted(ag.items(), key=lambda x: -x[1]["n"]):
        pct = round(g["n"] / len(regs) * 100)
        ultima = max(g["datas"])
        heat, rot = calor_de(ultima, ate)
        tag, ntag = g["tags"].most_common(1)[0]
        ws.append([tema, g["n"], barra(pct), f"{tag}  ({ntag/g['n']*100:.0f}%)",
                   media_fmt(g["total"]), media_fmt(g["tech"]),
                   date.fromisoformat(ultima).strftime("%d/%m/%Y"), rot])
        formatar(ws, ws.max_row, len(COLS), heat, pct, (2, 4, 5, 6), (7, 8))

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"
    ws.conditional_formatting.add(f"B2:B{ws.max_row}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="2E6B3E", showValue=True))
    return len(ag)


def aba_conversas(wb: Workbook, regs: list[dict], ate: str) -> int:
    ws = wb.create_sheet("Conversas")
    COLS = ["Tema", "Conversas", "Peso", "Mensagens (mediana)",
            "Duração (mediana)", "Última incidência", "Recência"]
    cabecalho(ws, COLS, [64, 11, 15, 19, 19, 14, 11])

    ag = collections.defaultdict(lambda: {"n": 0, "datas": [], "msg": [], "dur": []})
    for r in regs:
        g = ag[r["tema"]]
        g["n"] += 1
        g["datas"].append(r["data"])
        if r.get("mensagens"):
            g["msg"].append(int(r["mensagens"]))
        # Duração vem do Bitrix como a sessão inteira, e a sessão só fecha quando o
        # atendente encerra — há conversas de 4 dias. Por isso mediana, não média: uma
        # sessão esquecida aberta no fim de semana desloca a média inteira.
        if r.get("duracao_min"):
            g["dur"].append(int(r["duracao_min"]))

    def mediana(v: list[int], suf: str) -> str:
        if not v:
            return "—"
        import statistics
        return f"{statistics.median(v):.0f} {suf}"

    for tema, g in sorted(ag.items(), key=lambda x: -x[1]["n"]):
        pct = round(g["n"] / len(regs) * 100)
        ultima = max(g["datas"])
        heat, rot = calor_de(ultima, ate)
        ws.append([tema, g["n"], barra(pct), mediana(g["msg"], "msg"),
                   mediana(g["dur"], "min"),
                   date.fromisoformat(ultima).strftime("%d/%m/%Y"), rot])
        formatar(ws, ws.max_row, len(COLS), heat, pct, (2, 4, 5), (6, 7))

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    ws.conditional_formatting.add(f"B2:B{ws.max_row}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="6B4E8F", showValue=True))
    return len(ag)


def legenda(wb: Workbook, dados: dict, de: str, ate: str):
    ws = wb.create_sheet("Legenda")
    tx_ch = ler(STORE / "taxonomia_nao_fiscal.json")
    tx_cv = ler(STORE / "taxonomia_nao_fiscal_conv.json")
    nch, ncv = len(dados["chamado"]), len(dados["conversa"])
    itens = [
        ("Escopo", "Somente registros NÃO fiscais. O que é fiscal está em "
                   "relatorio_executivo_<período>.xlsx, com estrutura de frente/subgrupo "
                   "própria. Os dois relatórios não se somam: são taxonomias distintas."),
        ("Período", f"{de} a {ate}. Janela pura: contagens, medianas, última incidência e "
                    "recência consideram apenas registros dentro do período. A régua de "
                    "recência é relativa ao fim da janela, não à data de execução."),
        ("Volume no período", f"{nch} chamados e {ncv} conversas não-fiscais."),
        ("CONFIABILIDADE — leia antes de decidir",
         "Estas duas taxonomias NÃO passaram por auditoria cega. A auditoria de 31/08/2026 "
         "(erro por consenso 6,7%, kappa 0,891) mediu apenas a taxonomia FISCAL. Para o "
         "não-fiscal não há taxa de erro conhecida: o agrupamento serve para dimensionar e "
         "priorizar investigação, não para fechar decisão de roadmap sem auditar antes."),
        ("Duas abas, não uma",
         f"A taxonomia de chamado tem {len(tx_ch.get('temas', []))} temas e a de conversa "
         f"{len(tx_cv.get('temas', []))}, derivadas separadamente, sem um único nome em "
         "comum. Somar chamado e conversa numa linha só daria um total sem referente — por "
         "isso cada canal tem sua aba e seu peso calculado sobre o próprio total."),
        ("Peso", "Participação do tema no total do seu canal. Compare temas dentro da mesma "
                 "aba; não compare entre abas."),
        ("Categoria predominante (aba Chamados)",
         "Tag que o Bitrix já grava entre colchetes no título ([Melhoria], [Erro], "
         "[Orçamento]...) mais frequente dentro do tema, e o quanto ela representa. É a "
         "classificação que o próprio suporte fez na abertura, útil para conferir se o tema "
         "da IA bate com o rótulo humano — divergência grande é sinal de fronteira ruim."),
        ("Tempo total do chamado  (createdDate → closedDate)",
         "Média de dias corridos entre abertura e fechamento, só dos concluídos. O número "
         "entre parênteses é quantos entraram na média. É tempo de calendário, não esforço."),
        ("Tempo tech (desenvolvimento)  (dateStart → changedDate)",
         "Base menor que a coluna anterior porque dateStart falta em boa parte dos chamados. "
         "NÃO subtraia uma da outra para obter fila — são conjuntos distintos."),
        ("Mensagens / Duração (aba Conversas)",
         "Mediana, não média: a duração do Bitrix mede a sessão até o atendente encerrá-la, "
         "e sessões esquecidas abertas por dias deslocariam a média inteira."),
        ("Última incidência / Recência",
         "Data do registro mais recente do tema no período. Vermelho: últimos 30 dias antes "
         "do fim da janela. Laranja: 30 a 60. Amarelo: mais de 60."),
        ("(sem classificação)",
         "Registro não fiscal que ainda não passou pelo classificador. Linha própria em vez "
         "de sumir da contagem."),
        ("Consistência", f"Taxonomias {tx_ch.get('versao', '?')} (chamado) e "
                         f"{tx_cv.get('versao', '?')} (conversa), com cache em "
                         "data/store/classificacao_nao_fiscal*.json. Registros já "
                         "classificados não são reprocessados: o mesmo período dá sempre o "
                         "mesmo resultado."),
    ]
    for k, v in itens:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 118
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=4, column=1).font = Font(bold=True, color="9C2119")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--de", required=True)
    ap.add_argument("--ate", default=date.today().isoformat())
    a = ap.parse_args()
    REPORT.mkdir(exist_ok=True)

    dados = carregar(a.de, a.ate)
    if not (dados["chamado"] or dados["conversa"]):
        print(f"Nenhum registro não-fiscal em {a.de}..{a.ate}.")
        return 1

    wb = Workbook()
    wb.remove(wb.active)
    n_ch = aba_chamados(wb, dados["chamado"], a.ate) if dados["chamado"] else 0
    n_cv = aba_conversas(wb, dados["conversa"], a.ate) if dados["conversa"] else 0
    legenda(wb, dados, a.de, a.ate)
    saida = REPORT / f"executivo_naofiscal_{a.de}_{a.ate}.xlsx"
    wb.save(saida)

    sem = sum(1 for v in dados.values() for r in v if r["tema"] == SEM_CLS)
    print(f"Janela {a.de} a {a.ate}")
    print(f"  {len(dados['chamado'])} chamados não-fiscais em {n_ch} temas")
    print(f"  {len(dados['conversa'])} conversas não-fiscais em {n_cv} temas")
    if sem:
        print(f"  {sem} ainda sem classificação")
    print("\n  Taxonomias não-fiscais nunca auditadas — ver aba Legenda.")
    print(f"\n  {saida.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
